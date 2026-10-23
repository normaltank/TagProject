"""Read-only, local audit of the supplied workbooks. Never saves an XLSX.

Outputs contain aggregates and source locators, not teacher biographies/emails.
Language-script counts are diagnostic features, never language identifications.
"""
from __future__ import annotations

import argparse
import collections
import csv
import hashlib
import json
import platform
import re
import statistics
import sys
import time
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlsplit

import openpyxl

SCRIPT_VERSION = "1.1.0"
ZH = re.compile(r"[\u3400-\u4dbf\u4e00-\u9fff\U00020000-\U0003134f]")
LATIN = re.compile(r"[A-Za-z\u00c0-\u024f]")
KANA = re.compile(r"[\u3040-\u30ff]")
CYRILLIC = re.compile(r"[\u0400-\u052f]")
HANGUL = re.compile(r"[\uac00-\ud7af\u1100-\u11ff]")
ARABIC = re.compile(r"[\u0600-\u06ff]")
SCRIPT_PATTERNS = {"han": ZH, "latin": LATIN, "kana": KANA, "cyrillic": CYRILLIC,
                   "hangul": HANGUL, "arabic": ARABIC,
                   "greek": re.compile(r"[\u0370-\u03ff\u1f00-\u1fff]"),
                   "hebrew": re.compile(r"[\u0590-\u05ff]"),
                   "devanagari": re.compile(r"[\u0900-\u097f]"),
                   "thai": re.compile(r"[\u0e00-\u0e7f]")}
KNOWN_SCRIPT_RANGES = re.compile("|".join(p.pattern for p in SCRIPT_PATTERNS.values()))
HTML = re.compile(r"</?[A-Za-z][^>]{0,120}>")
INTEGER = re.compile(r"^[1-9][0-9]*$")
RESEARCH_CUE = re.compile(r"研究方向|科研方向|研究领域|research interests|research areas", re.I)


def sha256(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def stringify(value):
    return "" if value is None else str(value)


def empty(value):
    return value is None or (isinstance(value, str) and not value.strip())


def parse_array(value):
    if empty(value):
        return "blank", []
    if not isinstance(value, str):
        return "non_string_cell", []
    try:
        parsed = json.loads(value)
    except (ValueError, TypeError):
        return "invalid_json", []
    if not isinstance(parsed, list):
        return "json_not_array", []
    if any(not isinstance(s, str) or not s.strip() for s in parsed):
        return "array_with_nonstring_or_empty", parsed
    return ("empty_array" if not parsed else "string_array"), parsed


def norm(text):
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", text)).strip()


def distribution(values):
    if not values:
        return {"n": 0}
    ordered = sorted(values)
    def q(p):
        position = (len(ordered) - 1) * p
        lo = int(position)
        hi = min(lo + 1, len(ordered) - 1)
        return round(ordered[lo] + (ordered[hi] - ordered[lo]) * (position - lo), 2)
    return {"n": len(values), "sum": sum(values), "min": ordered[0], "p25": q(.25), "p50": q(.5),
            "p75": q(.75), "p90": q(.9), "p95": q(.95), "p99": q(.99),
            "max": ordered[-1], "mean": round(statistics.mean(values), 2),
            "quantile_method": "linear interpolation at (n-1)*p"}


def file_info(path):
    stat = path.stat()
    return {"path": str(path.resolve()), "bytes": stat.st_size,
            "modified_utc": datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(),
            "sha256_before": sha256(path)}


def write_json(path, obj):
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")


def write_csv(path, rows, fields):
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def audit_vocabulary(path, out):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    if len(wb.sheetnames) != 1:
        raise ValueError("This audit expects one vocabulary sheet; inspect new sheets before continuing")
    ws = wb.worksheets[0]
    row_iter = ws.iter_rows(values_only=True)
    headers = list(next(row_iter))
    required = {"id", "family_tag", "primary_tag", "secondary_tag", "deleted"}
    if not required.issubset(headers):
        raise ValueError(f"Unexpected vocabulary headers: {headers}")
    columns = {h: openpyxl.utils.get_column_letter(i + 1) for i, h in enumerate(headers)}
    rows, occurrences, problems = [], [], []
    names = {"L1": collections.defaultdict(list), "L2": collections.defaultdict(list), "L3": collections.defaultdict(list)}
    field_missing = collections.Counter()
    ids, parse_counts, deleted = collections.defaultdict(list), collections.Counter(), collections.Counter()
    for row_number, values in enumerate(row_iter, 2):
        if all(empty(v) for v in values):
            continue
        record = dict(zip(headers, values))
        for field, value in record.items():
            if empty(value):
                field_missing[field] += 1
        ids[stringify(record["id"])].append(row_number)
        deleted[stringify(record["deleted"])] += 1
        l1, l2 = stringify(record["family_tag"]), stringify(record["primary_tag"])
        state, children = parse_array(record["secondary_tag"])
        parse_counts[state] += 1
        row_data = {"sheet": ws.title, "excel_row": row_number, "source_id": stringify(record["id"]),
                    "family_tag": l1, "primary_tag": l2, "secondary_status": state,
                    "secondary_count": len(children), "deleted_raw": stringify(record["deleted"])}
        rows.append(row_data)
        names["L1"][l1].append({"row": row_number, "field": "family_tag", "cell": f"{columns['family_tag']}{row_number}"})
        names["L2"][l2].append({"row": row_number, "field": "primary_tag", "cell": f"{columns['primary_tag']}{row_number}", "family": l1})
        for child_index, child in enumerate(children):
            if not isinstance(child, str):
                continue
            occurrence = {"occurrence_id": f"r{row_number}:secondary_tag:{child_index}",
                          "sheet": ws.title, "excel_row": row_number, "cell": f"{columns['secondary_tag']}{row_number}",
                          "json_index_zero_based": child_index, "source_row_id": stringify(record["id"]),
                          "family_tag": l1, "primary_tag": l2, "label": child,
                          "normalized_for_diagnostic_only": norm(child)}
            occurrences.append(occurrence)
            names["L3"][child].append(occurrence)
            if child.endswith(("与", "和", "及", "、", "（", "(", "：", ":")) or child.count("（") != child.count("）") or child.count("(") != child.count(")"):
                problems.append({"issue": "possible_truncation_or_unbalanced_name", **occurrence})
        if state not in {"string_array", "empty_array", "blank"}:
            problems.append({"issue": "array_parse_problem", **row_data})
    duplicate_l3 = {k: v for k, v in names["L3"].items() if len(v) > 1}
    normalized = collections.defaultdict(set)
    for label in names["L3"]:
        normalized[norm(label)].add(label)
    crosslevel = {label: sorted(level for level in names if label in names[level])
                  for label in set().union(*(set(group) for group in names.values()))}
    crosslevel = {k: v for k, v in crosslevel.items() if len(v) > 1}
    summary = {
        "sheet": ws.title, "declared_max_row": ws.max_row, "declared_max_column": ws.max_column,
        "headers": headers, "columns": columns, "nonempty_data_rows": len(rows),
        "missing_fields": dict(field_missing), "deleted_distribution_raw": dict(deleted),
        "source_row_id_unique_count": len(ids), "duplicate_source_row_ids": {k: v for k, v in ids.items() if len(v) > 1},
        "hierarchy_convention": "L1=family_tag; L2=primary_tag; L3=elements of secondary_tag JSON array; this is a source interpretation, not a validated ontology",
        "l1_unique_exact_strings": len(names["L1"]), "l2_unique_exact_strings": len(names["L2"]),
        "l3_occurrences": len(occurrences), "l3_unique_exact_strings": len(names["L3"]),
        "l3_duplicate_string_groups": len(duplicate_l3), "l3_duplicate_extra_occurrences": sum(len(v) - 1 for v in duplicate_l3.values()),
        "l3_duplicates": duplicate_l3,
        "l3_diagnostic_normalization_collisions": {k: sorted(v) for k, v in normalized.items() if len(v) > 1},
        "crosslevel_exact_strings": crosslevel, "array_parse_status": dict(parse_counts),
        "rows_without_l3": [r for r in rows if r["secondary_count"] == 0],
        "name_issues_heuristic_only": problems,
        "queries": {term: {level: entries.get(term, []) for level, entries in names.items()}
                    for term in ("分离工程", "膜分离", "膜分离技术")},
        "formal_concept_ids_present": False,
        "formal_concept_id_note": "id identifies a source L2 row; no separate L1/L3 stable concept IDs or definitions are supplied"}
    write_csv(out / "vocabulary_rows.csv", rows, list(rows[0]))
    write_csv(out / "vocabulary_occurrences.csv", occurrences, list(occurrences[0]))
    wb.close()
    return summary, names


def audit_teachers(path, out, vocab_names):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    if len(wb.sheetnames) != 1:
        raise ValueError("This audit expects one teacher sheet; inspect new sheets before continuing")
    ws = wb.worksheets[0]
    row_iter = ws.iter_rows(values_only=True)
    headers = list(next(row_iter))
    required = {"id", "name", "intro", "tags", "homepage", "deleted"}
    if not required.issubset(headers):
        raise ValueError(f"Unexpected teacher headers: {headers}")
    columns = {h: openpyxl.utils.get_column_letter(i + 1) for i, h in enumerate(headers)}
    counts, missing, field_types = collections.Counter(), collections.Counter(), collections.defaultdict(collections.Counter)
    distributions = {h: collections.Counter() for h in ("deleted", "enrichment_status", "is_interdisciplinary", "aminer_synced")}
    ids, intro_hashes, homepage_rows = collections.defaultdict(list), collections.defaultdict(list), collections.defaultdict(list)
    blank_rows, malformed, lengths, all_lengths, tag_lengths = [], [], [], [], []
    script_totals, script_records, profiles = collections.Counter(), collections.Counter(), collections.Counter()
    tag_status, tag_counts, tag_levels = collections.Counter(), collections.Counter(), collections.Counter()
    unmatched, examples = collections.defaultdict(list), collections.defaultdict(list)
    tag_duplicate_rows, formula_cells, text_flags = [], [], collections.Counter()
    domain_counts = collections.Counter()
    tag_rows_by_string = collections.defaultdict(list)
    subsets = collections.defaultdict(lambda: {"counts": collections.Counter(), "missing": collections.Counter(),
                                                "lengths": [], "tag_status": collections.Counter(),
                                                "script_profiles": collections.Counter(), "length_bins": collections.Counter()})
    unknown_letters = collections.Counter()
    raw_nonempty_missing = collections.Counter()
    row_fields = ["sheet", "excel_row", "teacher_id", "intro_chars", "intro_sha256"] + [s + "_chars" for s in SCRIPT_PATTERNS] + ["other_letter_chars", "detected_script_ranges", "script_profile", "language_verification", "research_cue", "html_like", "replacement_char", "homepage_domain", "tags_status", "tags_count", "deleted_raw", "enrichment_status_raw"]
    tag_fields = ["sheet", "excel_row", "teacher_id", "cell", "json_index_zero_based", "historical_label", "matching_levels", "matching_occurrences", "match_status"]
    with (out / "teacher_rows.csv").open("w", newline="", encoding="utf-8-sig") as rf, (out / "historical_assignments.csv").open("w", newline="", encoding="utf-8-sig") as tf:
        row_writer, tag_writer = csv.DictWriter(rf, fieldnames=row_fields), csv.DictWriter(tf, fieldnames=tag_fields)
        row_writer.writeheader()
        tag_writer.writeheader()
        for row_number, values in enumerate(row_iter, 2):
            if all(empty(v) for v in values):
                blank_rows.append(row_number)
                continue
            counts["nonempty_data_rows"] += 1
            record = dict(zip(headers, values))
            for field, value in record.items():
                if empty(value):
                    raw_nonempty_missing[field] += 1
            raw_id = stringify(record["id"])
            valid_id = bool(INTEGER.fullmatch(raw_id))
            if not valid_id:
                counts["nonempty_rows_without_valid_positive_integer_id"] += 1
                malformed.append({"sheet": ws.title, "excel_row": row_number, "id_preview": raw_id[:80],
                                  "populated_fields": [h for h, v in record.items() if not empty(v)],
                                  "note": "non-record candidate; retained in original, excluded from teacher denominators"})
                continue
            counts["valid_id_rows"] += 1
            subset = subsets[stringify(record["deleted"])]
            subset["counts"]["records"] += 1
            ids[raw_id].append(row_number)
            for field, value in record.items():
                if empty(value):
                    missing[field] += 1
                    subset["missing"][field] += 1
                field_types[field][type(value).__name__] += 1
                if isinstance(value, str) and value.startswith("="):
                    formula_cells.append({"row": row_number, "field": field, "cell": f"{columns[field]}{row_number}"})
            for field in distributions:
                distributions[field][stringify(record.get(field))] += 1
            intro = stringify(record["intro"])
            length = len(intro)
            all_lengths.append(length)
            subset["lengths"].append(length)
            length_bin = ("0" if length == 0 else "1-99" if length < 100 else "100-999" if length < 1000
                          else "1000-3999" if length < 4000 else "4000-7999" if length < 8000
                          else "8000-15999" if length < 16000 else "16000-32766" if length < 32767
                          else "32767" if length == 32767 else "over_32767")
            subset["length_bins"][length_bin] += 1
            intro_hash = hashlib.sha256(intro.encode("utf-8")).hexdigest() if intro.strip() else ""
            if intro.strip():
                counts["nonblank_intro_rows"] += 1
                lengths.append(length)
                intro_hashes[intro_hash].append(row_number)
            else:
                counts["blank_intro_rows"] += 1
                subset["counts"]["blank_intro_rows"] += 1
            measures = {script: len(pattern.findall(intro)) for script, pattern in SCRIPT_PATTERNS.items()}
            residual_letters = collections.Counter(c for c in KNOWN_SCRIPT_RANGES.sub("", intro) if c.isalpha())
            unknown_letters.update(residual_letters)
            measures["other_letter"] = sum(residual_letters.values())
            for script, amount in measures.items():
                script_totals[script] += amount
                if amount:
                    script_records[script] += 1
            if not intro.strip():
                profile = "blank"
            elif measures["han"] and measures["latin"]:
                profile = "Han+Latin"
            elif measures["han"]:
                profile = "Han_without_Latin"
            elif measures["latin"]:
                profile = "Latin_without_Han"
            else:
                profile = "other_or_nonletter"
            profiles[profile] += 1
            subset["script_profiles"][profile] += 1
            html_like, replacement, cue = bool(HTML.search(intro)), "\ufffd" in intro, bool(RESEARCH_CUE.search(intro))
            for flag, condition in (("html_like", html_like), ("replacement_character", replacement), ("explicit_research_cue", cue), ("intro_over_8000_chars", length > 8000), ("intro_over_32000_chars", length > 32000), ("intro_32767_chars", length == 32767), ("intro_under_100_nonblank_chars", 0 < length < 100)):
                if condition:
                    text_flags[flag] += 1
                    subset["counts"][flag] += 1
                    if len(examples[flag]) < 20:
                        examples[flag].append(row_number)
            for script in list(SCRIPT_PATTERNS) + ["other_letter"]:
                if script in {"han", "latin"}:
                    continue
                if measures[script] and len(examples[script + "_any"]) < 25:
                    examples[script + "_any"].append(row_number)
                if measures[script] >= 10 and len(examples[script + "_at_least_10"]) < 25:
                    examples[script + "_at_least_10"].append(row_number)
            if profile == "Latin_without_Han" and length >= 200 and len(examples["latin_no_han_at_least_200"]) < 25:
                examples["latin_no_han_at_least_200"].append(row_number)
            if length > 0 and not cue and len(examples["no_research_heading"]) < 10:
                examples["no_research_heading"].append(row_number)
            homepage = stringify(record["homepage"]).strip()
            domain = ""
            if homepage:
                homepage_rows[homepage].append(row_number)
                try:
                    domain = (urlsplit(homepage).hostname or "").lower()
                except ValueError:
                    pass
                domain_counts[domain or "unparsed"] += 1
            state, tags = parse_array(record["tags"])
            tag_status[state] += 1
            subset["tag_status"][state] += 1
            string_tags = [t for t in tags if isinstance(t, str) and t.strip()]
            subset["counts"]["historical_assignment_occurrences"] += len(string_tags)
            tag_lengths.append(len(string_tags))
            if string_tags:
                counts["rows_with_historical_tags"] += 1
                subset["counts"]["rows_with_historical_tags"] += 1
            if len(string_tags) != len(set(string_tags)):
                tag_duplicate_rows.append(row_number)
            if not intro.strip() and string_tags:
                counts["blank_intro_rows_with_historical_tags"] += 1
                subset["counts"]["blank_intro_rows_with_historical_tags"] += 1
                if len(examples["blank_intro_with_tags"]) < 20:
                    examples["blank_intro_with_tags"].append(row_number)
            for i, tag in enumerate(tags):
                if not isinstance(tag, str) or not tag.strip():
                    continue
                tag_counts[tag] += 1
                if len(tag_rows_by_string[tag]) < 20:
                    tag_rows_by_string[tag].append(row_number)
                levels = [level for level in vocab_names if tag in vocab_names[level]]
                matches = sum(len(vocab_names[level].get(tag, [])) for level in levels)
                level_key = "+".join(levels) if levels else "unmatched"
                tag_levels[level_key] += 1
                if not levels:
                    unmatched[tag].append(row_number)
                tag_writer.writerow({"sheet": ws.title, "excel_row": row_number, "teacher_id": raw_id,
                                     "cell": f"{columns['tags']}{row_number}", "json_index_zero_based": i,
                                     "historical_label": tag, "matching_levels": level_key,
                                     "matching_occurrences": matches, "match_status": "exact_string_only_not_semantic_validation" if levels else "no_exact_string_match"})
            row_writer.writerow({"sheet": ws.title, "excel_row": row_number, "teacher_id": raw_id,
                                 "intro_chars": length, "intro_sha256": intro_hash,
                                 **{k + "_chars": v for k, v in measures.items()}, "script_profile": profile,
                                 "detected_script_ranges": "+".join(s for s, amount in measures.items() if amount),
                                 "language_verification": "not_human_verified",
                                 "research_cue": int(cue), "html_like": int(html_like), "replacement_char": int(replacement),
                                 "homepage_domain": domain, "tags_status": state, "tags_count": len(string_tags),
                                 "deleted_raw": stringify(record["deleted"]), "enrichment_status_raw": stringify(record.get("enrichment_status"))})
            if counts["valid_id_rows"] % 20000 == 0:
                print(f"Audited {counts['valid_id_rows']} valid-ID rows", flush=True)
    duplicate_intros = [{"intro_sha256": h, "rows": r} for h, r in intro_hashes.items() if len(r) > 1]
    duplicates_homepage = [{"homepage_sha256": hashlib.sha256(h.encode()).hexdigest(), "rows": r} for h, r in homepage_rows.items() if len(r) > 1]
    counts["blank_physical_rows"] = len(blank_rows)
    counts["unique_valid_ids"] = len(ids)
    counts["duplicate_valid_id_groups"] = sum(len(r) > 1 for r in ids.values())
    counts["historical_assignment_occurrences"] = sum(tag_counts.values())
    counts["historical_unique_exact_strings"] = len(tag_counts)
    counts["duplicate_intro_groups"] = len(duplicate_intros)
    counts["rows_in_duplicate_intro_groups"] = sum(len(r["rows"]) for r in duplicate_intros)
    counts["duplicate_homepage_groups"] = len(duplicates_homepage)
    counts["rows_in_duplicate_homepage_groups"] = sum(len(r["rows"]) for r in duplicates_homepage)
    summary = {"sheet": ws.title, "declared_max_row": ws.max_row, "declared_max_column": ws.max_column,
               "headers": headers, "columns": columns, "counts": dict(counts),
               "teacher_denominator": "rows with a positive integer in id; no deleted/status filter; uniqueness is reported separately",
               "missing_field_counts_valid_id_rows": dict(missing), "field_python_types": {k: dict(v) for k, v in field_types.items()},
               "missing_field_counts_all_nonempty_rows_not_teacher_denominator": dict(raw_nonempty_missing),
               "by_deleted_raw": {key: {"denominator": "positive-integer-ID rows with this exact deleted value; business meaning unverified",
                                         "counts": dict(value["counts"]), "missing_fields": dict(value["missing"]),
                                         "intro_characters": distribution(value["lengths"]), "length_bins_chars": dict(value["length_bins"]),
                                         "tag_parse_status": dict(value["tag_status"]), "script_profiles": dict(value["script_profiles"])}
                                  for key, value in subsets.items()},
               "raw_status_distributions": {k: dict(v) for k, v in distributions.items()},
               "blank_excel_rows": blank_rows, "nonrecord_candidates": malformed,
               "duplicate_ids": {k: v for k, v in ids.items() if len(v) > 1},
               "intro_characters_all_valid_id_rows": distribution(all_lengths),
               "intro_characters_nonblank_only": distribution(lengths),
               "historical_tags_per_valid_id_row": distribution(tag_lengths),
               "intro_script_character_totals": dict(script_totals), "intro_rows_containing_script": dict(script_records),
               "intro_script_profiles": dict(profiles),
               "other_letter_characters": [{"character": c, "unicode_name": unicodedata.name(c, "UNKNOWN"), "count": n} for c, n in unknown_letters.most_common()],
               "language_caveat": "Unicode script is not language. English, French, German etc share Latin; Han may occur in Japanese. Presence may be a citation/acronym/name; not a prose-language prevalence estimate.",
               "text_quality_flags_heuristic": dict(text_flags), "sample_rows_for_local_inspection": dict(examples),
               "formula_like_cells": formula_cells,
               "tag_parse_status": dict(tag_status), "historical_exact_match_levels": dict(tag_levels),
               "historical_unmatched_strings": {k: {"count": len(v), "rows": v} for k, v in unmatched.items()},
               "historical_duplicate_tag_array_rows": tag_duplicate_rows,
               "historical_tag_frequencies": [{"label": k, "occurrences": v, "first_rows": tag_rows_by_string[k]} for k, v in tag_counts.most_common()],
               "homepage_domain_top30": domain_counts.most_common(30),
               "duplicate_groups_file": "duplicate_groups.json"}
    write_json(out / "duplicate_groups.json", {"exact_nonblank_intro": duplicate_intros, "exact_nonblank_homepage": duplicates_homepage})
    wb.close()
    return summary


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--teachers", type=Path, required=True)
    parser.add_argument("--vocabulary", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True, help="A new output directory; existing files are never overwritten")
    args = parser.parse_args()
    if args.out.exists() and any(args.out.iterdir()):
        raise SystemExit("Output directory is nonempty; choose a new run directory")
    args.out.mkdir(parents=True, exist_ok=True)
    start = time.perf_counter()
    manifest = {"run_started_utc": datetime.now(timezone.utc).isoformat(), "script_version": SCRIPT_VERSION,
                "script_sha256": sha256(Path(__file__)), "python": sys.version, "executable": sys.executable,
                "openpyxl_version": openpyxl.__version__, "platform": platform.platform(),
                "argv": sys.argv, "read_only": True,
                "inputs": {"teachers": file_info(args.teachers), "vocabulary": file_info(args.vocabulary)}}
    (args.out / "audit_script_snapshot.py").write_bytes(Path(__file__).read_bytes())
    write_json(args.out / "manifest.json", manifest)
    vocabulary, names = audit_vocabulary(args.vocabulary, args.out)
    teachers = audit_teachers(args.teachers, args.out, names)
    for info in manifest["inputs"].values():
        info["sha256_after"] = sha256(Path(info["path"]))
        info["unchanged"] = info["sha256_before"] == info["sha256_after"]
        if not info["unchanged"]:
            raise RuntimeError(f"Input changed during audit: {info['path']}")
    manifest["run_finished_utc"] = datetime.now(timezone.utc).isoformat()
    manifest["wall_seconds"] = round(time.perf_counter() - start, 3)
    manifest["output_files"] = [p.name for p in args.out.iterdir()] + ["audit_summary.json"]
    write_json(args.out / "audit_summary.json", {"vocabulary": vocabulary, "teachers": teachers})
    write_json(args.out / "manifest.json", manifest)
    print(json.dumps({"output": str(args.out.resolve()), "wall_seconds": manifest["wall_seconds"],
                      "teacher_counts": teachers["counts"],
                      "vocabulary_counts": {k: vocabulary[k] for k in ("l1_unique_exact_strings", "l2_unique_exact_strings", "l3_occurrences", "l3_unique_exact_strings")},
                      "inputs_unchanged": all(v["unchanged"] for v in manifest["inputs"].values())}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
