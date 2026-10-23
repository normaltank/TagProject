"""Render/validate Agent-authored examples. Does NOT infer tags for a new biography."""
from __future__ import annotations
import argparse
import collections
import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path
import openpyxl

VERSION = "0.1.1"
CONTACT = re.compile(r"[\w.+%-]+@[\w.-]+\.[A-Za-z]{2,}|(?<!\d)1[3-9]\d{9}(?!\d)|(?<!\d)0\d{2,3}[- ]?\d{7,8}(?!\d)")
POSITIVE_ROLES = {"explicit_research_direction", "explicit_research_content", "explicit_research_interest"}
STATE_ZH = {"supported_pending_review": "有原文支持，标签建议待真人复核", "insufficient_no_intro": "无简介，资料不足不赋标",
            "source_only_administrative": "行政岗位／菜单，暂不赋标", "source_only_article": "文章页与教师关系未确认，暂不赋标",
            "unreadable_text": "乱码，无法可靠判断", "invalid_numeric_cell": "数字型简介，待修复不赋标",
            "invalid_date_cell": "日期型简介，待修复不赋标", "formula_not_research_text": "公式型简介，待修复不赋标"}


def sha(path):
    with path.open("rb") as f:
        return hashlib.file_digest(f, "sha256").hexdigest()


def save_json(path, value):
    if path.exists():
        raise RuntimeError(f"Will not overwrite {path}")
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--teachers", type=Path, required=True)
    ap.add_argument("--vocabulary", type=Path, required=True)
    ap.add_argument("--audit-manifest", type=Path, required=True)
    ap.add_argument("--draft", type=Path, required=True)
    ap.add_argument("--mapping", type=Path, required=True)
    ap.add_argument("--out", type=Path, required=True)
    a = ap.parse_args()
    a.out.mkdir(parents=True, exist_ok=True)
    targets = [a.out / x for x in ("examples.json", "REVIEW.md", "checks.json", "manifest.json", "human_feedback.json")]
    assert not any(p.exists() for p in targets), "Choose a new output directory for a new result version"
    baseline = json.loads(a.audit_manifest.read_text(encoding="utf-8"))
    hashes = {"teachers": sha(a.teachers), "vocabulary": sha(a.vocabulary)}
    assert all(hashes[k] == baseline["inputs"][k]["sha256_before"] for k in hashes)
    draft = json.loads(a.draft.read_text(encoding="utf-8"))
    mapping = json.loads(a.mapping.read_text(encoding="utf-8"))
    selected = {c["excel_row"]: c for c in draft["cases"]}
    decisions = {c["excel_row"]: c for c in mapping["mappings"]}
    assert len(selected) == draft["case_count"] == 16
    vocabulary_wb = openpyxl.load_workbook(a.vocabulary, read_only=True, data_only=False)
    vs = vocabulary_wb["review_tagscollection_20261022"]
    rows = list(vs.iter_rows(values_only=True))
    header = {name: i for i, name in enumerate(rows[0])}
    registry = collections.defaultdict(list)
    all_names = set()
    for rn, row in enumerate(rows[1:], 2):
        record = {key: row[index] for key, index in header.items()}
        family, primary = record["family_tag"], record["primary_tag"]
        all_names.update((family, primary))
        common = {"file": "Sheets/review_tagscollection_20261022.xlsx", "file_sha256": hashes["vocabulary"],
                  "sheet": vs.title, "excel_row": rn, "source_row_id": str(record["id"]),
                  "formal_concept_id": None}
        registry[("L2", primary)].append({**common, "level": "L2", "field": "primary_tag", "cell": f"E{rn}",
                                          "json_index_zero_based": None, "source_occurrence_id": f"r{rn}:primary_tag",
                                          "label": primary, "path": [family, primary]})
        for index, label in enumerate(json.loads(record["secondary_tag"] or "[]")):
            all_names.add(label)
            registry[("L3", label)].append({**common, "level": "L3", "field": "secondary_tag", "cell": f"F{rn}",
                                            "json_index_zero_based": index,
                                            "source_occurrence_id": f"r{rn}:secondary_tag:{index}",
                                            "label": label, "path": [family, primary, label]})
    vocabulary_wb.close()
    teacher_wb = openpyxl.load_workbook(a.teachers, read_only=True, data_only=False)
    ts = teacher_wb["teacher_20261022"]
    cases, verification, name_tokens = [], [], []
    for rn, cells in enumerate(ts.iter_rows(max_row=max(selected), max_col=15), 1):
        if rn not in selected:
            continue
        c = json.loads(json.dumps(selected[rn], ensure_ascii=False))
        teacher_id = str(cells[0].value)
        assert teacher_id == c.pop("expected_teacher_id"), (rn, teacher_id)
        own_name = str(cells[1].value or "").strip()
        if own_name:
            name_tokens.append(own_name)
        raw = cells[14].value
        text = raw if isinstance(raw, str) else ""
        evidence = []
        for item in c["evidence"]:
            quote = item["quote"]
            anchor = item.get("after")
            anchor_start = text.find(anchor) if anchor else 0
            assert anchor_start >= 0, (rn, "missing anchor", anchor)
            start = text.find(quote, anchor_start + (len(anchor) if anchor else 0))
            assert start >= 0, (rn, "missing quotation", quote)
            end = start + len(quote)
            assert text[start:end] == quote
            assert not CONTACT.search(quote), (rn, "contact in excerpt")
            assert not own_name or own_name not in quote, (rn, "unneeded person name in excerpt")
            positions = [m.start() for m in re.finditer(re.escape(quote), text)]
            evidence.append({**item, "source_cell": f"O{rn}", "start_zero_based": start, "end_exclusive": end,
                             "identical_quote_occurrences_in_cell": len(positions),
                             "chosen_occurrence_one_based": positions.index(start) + 1,
                             "quote_exactly_verified": True})
            verification.append({"teacher_id": teacher_id, "excel_row": rn, "evidence_id": item["evidence_id"],
                                 "start_zero_based": start, "end_exclusive": end, "exact_match": True})
        c["evidence"] = evidence
        evidence_by_id = {e["evidence_id"]: e for e in evidence}
        direction_ids = {d["direction_id"] for d in c["directions"]}
        assert len(evidence_by_id) == len(evidence) and len(direction_ids) == len(c["directions"])
        for d in c["directions"]:
            assert d["evidence_ids"]
            assert all(evidence_by_id[x]["role"] in POSITIVE_ROLES for x in d["evidence_ids"])
        def resolve(items, status):
            result = []
            for item in items:
                key = (item.get("level", "L3"), item["label"])
                matches = registry[key]
                assert len(matches) == 1, (rn, "ambiguous/missing vocabulary reference", key)
                assert all(x in direction_ids for x in item["direction_ids"])
                assert all(x in evidence_by_id for x in item["evidence_ids"])
                if status == "recommended_pending_human_review":
                    assert item["direction_ids"] and item["evidence_ids"]
                    assert all(evidence_by_id[x]["role"] in POSITIVE_ROLES for x in item["evidence_ids"])
                certainty = ("近似映射、范围待审" if item["relation"] == "broader_approximation_scope_review" else
                             "候选边界或证据不足，暂不赋标" if status == "review_only_not_assigned" else
                             "语义近合建议，待真人确认" if item["relation"] == "close_existing_concept" else
                             "上位／较粗映射建议，待真人确认")
                result.append({**item, "status": status, "mapping_certainty_zh": certainty,
                               "vocabulary_source": matches[0], "formal_assignment_made": False})
            return result
        m = decisions.get(rn, {"recommended": [], "review_only": [], "gaps": []})
        c["tag_recommendations"] = resolve(m["recommended"], "recommended_pending_human_review")
        c["existing_candidates_not_assigned"] = resolve(m["review_only"], "review_only_not_assigned")
        c["coverage_and_gap_candidates"] = [{**g, "semantic_gap_confirmed": False,
            "exact_label_name_lookup": [{"text": term, "exact_name_present_in_any_level": term in all_names}
                                        for term in re.split(r"／", g["candidate_description"])],
            "formal_label_created": False} for g in m["gaps"]]
        c["teacher_id"] = teacher_id
        c["source"] = {"file": "Sheets/teacher_20261022.xlsx", "file_sha256": hashes["teachers"], "sheet": ts.title,
                        "excel_row": rn, "field": "intro", "cell": f"O{rn}", "id_cell": f"A{rn}",
                        "deleted_raw": cells[11].value, "python_value_type": type(raw).__name__, "excel_cell_type": cells[14].data_type,
                        "intro_text_chars": len(text) if isinstance(raw, str) else None,
                        "intro_text_sha256": hashlib.sha256(text.encode("utf-8")).hexdigest() if text else None,
                        "scalar_display_if_not_text": str(raw) if raw is not None and not isinstance(raw, str) else None,
                        "text_offsets_applicable": isinstance(raw, str), "empty_cell": raw is None}
        c["source_review_scope"] = "selected evidence for a development example; not exhaustive of all research directions or all source quality issues"
        c["historical_labels_supplied_this_round"] = False
        c["human_review_status"] = "not_reviewed"
        if not c["directions"]:
            assert not c["tag_recommendations"]
        cases.append(c)
    teacher_wb.close()
    assert len(cases) == 16
    counts = {"cases": len(cases), "cases_with_tag_recommendations": sum(bool(c["tag_recommendations"]) for c in cases),
              "cases_without_tag_recommendations": sum(not c["tag_recommendations"] for c in cases),
              "source_content_directions": sum(len(c["directions"]) for c in cases),
              "tag_recommendation_occurrences": sum(len(c["tag_recommendations"]) for c in cases),
              "approximate_scope_review_recommendations": sum(t["relation"] == "broader_approximation_scope_review" for c in cases for t in c["tag_recommendations"]),
              "human_confirmed_mappings": 0,
              "candidate_only_occurrences_not_assigned": sum(len(c["existing_candidates_not_assigned"]) for c in cases),
              "exact_quotation_checks": len(verification)}
    data = {"schema_version": VERSION, "batch_id": a.out.name, "evaluator_type": "Agent",
            "purpose": "selected teacher development examples for human review; not automatic tagging, independent gold, or model performance",
            "prior_audit_context_exposure": True, "fully_blinded": False, "historical_label_comparison_run": False,
            "formal_concept_ids_created": False, "formal_assignments_written": False,
            "offset_convention": "Python Unicode code points, 0-based start inclusive and end exclusive, original cell value without normalization",
            "vocabulary_identity_note": "Source row IDs, hierarchy paths and array occurrences are locators, not newly minted formal concept IDs",
            "directions_draft_sha256": sha(a.draft), "mapping_decisions_sha256": sha(a.mapping), "counts": counts, "cases": cases}
    data_text = json.dumps(data, ensure_ascii=False)
    # Scan human-facing text, not SHA256 metadata that can coincidentally contain phone-shaped digits.
    human_text = []
    for c in cases:
        human_text += [c["rationale"]] + [e["quote"] for e in c["evidence"]] + [d["summary"] for d in c["directions"]]
        human_text += [t["reason"] for key in ("tag_recommendations", "existing_candidates_not_assigned") for t in c[key]]
        human_text += [g["candidate_description"] + " " + g["note"] for g in c["coverage_and_gap_candidates"]]
    assert all(not CONTACT.search(value) for value in human_text)
    assert all(name not in data_text for name in name_tokens if len(name) > 1)
    assert len(verification) == 40
    result_hashes = {"teachers": sha(a.teachers), "vocabulary": sha(a.vocabulary)}
    assert result_hashes == hashes
    checks = {"checked_date": "2026-10-23", "all_pass": True, "case_ids_unique": len({c["teacher_id"] for c in cases}) == 16,
              "counts": counts, "all_40_quotes_exact": len(verification) == 40 and all(v["exact_match"] for v in verification),
              "all_vocabulary_references_verified_against_original_workbook": True,
              "recommendations_use_explicit_research_evidence_only": True,
              "nonassigned_cases_have_no_tag_recommendations": True,
              "contact_and_selected_teacher_name_scan_passed": True,
              "historical_tags_column_not_consumed": True, "input_hashes_unchanged": result_hashes == hashes,
              "quote_checks": verification, "limits": "These checks validate locators, exact quotes and structural rules; semantic directions/mapping quality await human review."}
    save_json(a.out / "examples.json", data)
    save_json(a.out / "checks.json", checks)
    feedback = {"batch_id": data["batch_id"], "review_status": "not_started", "not_gold_before_review": True,
                "instructions": "核查原文方向与证据角色、标签是否过宽或越界、遗漏细节是否需要组合/别名或只留描述；不以旧标签作答案。修改后保留理由。",
                "cases": [{"teacher_id": c["teacher_id"], "excel_row": c["excel_row"], "status": "not_reviewed",
                           "direction_corrections": [], "accepted_source_tag_occurrences": [], "rejected_source_tag_occurrences": [],
                           "gap_disposition": [], "no_tag_route_acceptable": None, "reviewer_notes": "", "review_minutes": None}
                          for c in cases]}
    save_json(a.out / "human_feedback.json", feedback)
    md = ["# 教师方向与标签：16条开发示例审阅稿", "",
          "2026-10-23 · 当前只推进教师主线。以下为Agent选取证据、提炼方向和人工配置映射的开发示例，均待真人复核；不是独立金标、模型性能或自动打标签原型。部分行的旧标签曾在此前审计上下文出现，本轮材料和判断未提供旧标签，不能称完全盲审。", "",
          "长简介只选取了本批讨论的证据，未宣称穷尽全部方向。证据不足仅表示本次材料不足，不表示教师没有该专长。较粗映射不意味着教师覆盖该父类的全部研究内容。当前词库没有提供定义，标为范围待审的近似映射需要重点复核。", "",
          f"本批：16条，{counts['cases_with_tag_recommendations']}条有标签建议、{counts['cases_without_tag_recommendations']}条不赋标；{counts['source_content_directions']}个选取的内容方向、{counts['tag_recommendation_occurrences']}次标签建议（未展开祖先），其中{counts['approximate_scope_review_recommendations']}次为近似映射、范围待审，真人已确认映射0次；另有{counts['candidate_only_occurrences_not_assigned']}个仅待审候选；{counts['exact_quotation_checks']}条短引已与原文精确核对。这些是开发产物数量，不是效果指标。", "",
          "[结构化结果](examples.json) · [定位校验](checks.json) · [真人反馈模板](human_feedback.json)", "",
          "来源：Sheets/teacher_20261022.xlsx，sheet teacher_20261022，证据字段O=intro；词库为Sheets/review_tagscollection_20261022.xlsx，sheet同名。字符区间使用0起点、右端不含，基于原始单元格文本，未先归一化。数字／日期／空值没有可用研究文本，字符区间不适用；公式引文只展示输入异常。两个原件哈希及配置版本见manifest.json，未改原表或创建正式概念ID。", "",
          "## 真人优先确认", "",
          "1. 方向是否由引文支持，课程、菜单、旧项目或文章主题有无被误当稳定专长。",
          "2. 资源与环境化工等上位近似的范围是否合适；较粗分类是否丢失关键区分。",
          "3. 未覆盖描述应保留原文、作为既有概念组合／别名，还是值得后续提案；本批不自动新建。",
          "4. 七条不赋标的路由是否合理，哪些需要补原文或确认对象关系。", ""]
    def tag_line(t):
        s = t["vocabulary_source"]
        index = f"，数组索引{s['json_index_zero_based']}（0起点）" if s["level"] == "L3" else "，无数组索引"
        return f"**{t['mapping_certainty_zh']}** · {' → '.join(s['path'])}；{s['level']}，源行ID={s['source_row_id']}，{s['cell']}{index}；证据{','.join(t['evidence_ids'])}。{t['reason']}"
    for i, c in enumerate(cases, 1):
        s = c["source"]
        md += [f"## {i:02d} · 教师ID {c['teacher_id']} · Excel行{s['excel_row']}", "",
               f"来源：teacher_20261022!{s['cell']}；ID位于{s['id_cell']}；deleted原值={s['deleted_raw']}；单元格类型{s['python_value_type']}/{s['excel_cell_type']}。",
               f"判断：**{STATE_ZH[c['status']]}**。{c['rationale']}", ""]
        if s["scalar_display_if_not_text"] is not None:
            md += [f"非文本值展示：{s['scalar_display_if_not_text']}；无研究文本字符区间。", ""]
        if s["empty_cell"]:
            md += ["单元格为空，无可引用文本；证据位置保留到源单元格。", ""]
        if c["directions"]:
            md += ["选取的研究方向：", ""]
            md += [f"- {d['direction_id']}：{d['summary']}（证据{','.join(d['evidence_ids'])}）。" for d in c["directions"]]
            md += [""]
        if c["evidence"]:
            md += ["原文短引与角色：", ""]
            for e in c["evidence"]:
                md += [f"- {e['evidence_id']} · {s['cell']}[{e['start_zero_based']},{e['end_exclusive']}) · {e['role']}："]
                md += ["  > " + line for line in e["quote"].splitlines()]
                if e["identical_quote_occurrences_in_cell"] > 1:
                    md += [f"  同短引在该单元格共{e['identical_quote_occurrences_in_cell']}次，此处固定第{e['chosen_occurrence_one_based']}次；不把重复出现算独立来源。"]
            md += [""]
        if c["tag_recommendations"]:
            md += ["现有标签建议（均待真人复核，未正式写入）：", ""]
            md += ["- " + tag_line(t) for t in c["tag_recommendations"]]
            md += [""]
        else:
            md += ["标签建议：无。", ""]
        if c["existing_candidates_not_assigned"]:
            md += ["仅待审、不赋标的既有候选：", ""]
            md += ["- " + tag_line(t) for t in c["existing_candidates_not_assigned"]]
            md += [""]
        if c.get("withheld_directions"):
            md += ["暂不认定的方向：", ""]
            md += [f"- {d['summary']}：{d['reason']}" for d in c["withheld_directions"]]
            md += [""]
        if c["coverage_and_gap_candidates"]:
            md += ["细节保留／未覆盖候选（不是新正式标签，语义缺口未获专家确认）：", ""]
            md += [f"- {g['candidate_description']}：{g['note']}" for g in c["coverage_and_gap_candidates"]]
            md += [""]
    md += ["## 本轮实际完成与下一项编码工作", "",
           "已运行：Python/openpyxl读取16条教师A:O材料与现有词库；Agent先写方向草稿再写映射；生成／校验工具检查短引偏移、ID及词库行／数组位置，未读取本轮旧标签列、未下载或调用模型。", "",
           "未运行：真人复核、可接收新简介的自动规则基线、语义检索／重排、LLM、训练与规模化赋标。下一项编码工作是A规则／关键词（必要时BM25）＋证据位置＋候选标签，并用这批开发案例检查；具体B模型、阈值与硬件未定，C待收益／成本证据。其他支线保留文档，不阻塞教师工作。", ""]
    (a.out / "REVIEW.md").write_text("\n".join(md), encoding="utf-8")
    manifest = {"built_utc": datetime.now(timezone.utc).isoformat(), "script_version": VERSION, "script_sha256": sha(Path(__file__)),
                "python_openpyxl": {"openpyxl": openpyxl.__version__}, "inputs_sha256_before": hashes,
                "inputs_sha256_after": result_hashes, "directions_draft_sha256": sha(a.draft),
                "mapping_decisions_sha256": sha(a.mapping), "counts": counts,
                "workflow": "Agent-authored source-direction draft -> Agent-authored mapping decisions -> deterministic rendering and locator validation",
                "automatic_tag_inference_run": False, "historical_comparison_run": False, "human_review_run": False,
                "prior_audit_context_exposure": True, "no_new_formal_concept_ids": True,
                "outputs_sha256": {p.name: sha(p) for p in targets if p.exists()}}
    save_json(a.out / "manifest.json", manifest)
    print(json.dumps({"output": str(a.out), "counts": counts, "locators_and_hashes_passed": True, "model_or_automatic_tagger_run": False}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
